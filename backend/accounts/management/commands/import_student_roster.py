from __future__ import annotations

import re
from collections import defaultdict
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from academics.defaults import CLASS_LADDER
from academics.models import ClassArm, ClassLevel, StudentIdSequence
from accounts.models import AccountType, ParentProfile, StudentProfile, User
from accounts.passwords import DEFAULT_PASSWORD

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    raise CommandError(
        "openpyxl is required. Install it with: pip install openpyxl"
    ) from exc


DEFAULT_ROSTER_DIR = Path(__file__).resolve().parents[2] / "data" / "rosters"

# Legacy spreadsheet class labels → canonical class levels in this app.
CLASS_MAP = {
    "CRECHE": "Creche",
    "PRENURSERY": "Pre-Nursery",
    "PRE-NURSERY": "Pre-Nursery",
    "PRE NURSERY": "Pre-Nursery",
    "NURSERYPRE": "Pre-Nursery",
    "DAYCARE": "Pre-Nursery",
    "DAY CARE": "Pre-Nursery",
    "NURSERY1": "Nursery 1",
    "NURSERY 1": "Nursery 1",
    "NURSERY2": "Nursery 2",
    "NURSERY 2": "Nursery 2",
    # School ladder ends at Nursery 2; treat Nursery 3 as Nursery 2.
    "NURSERY3": "Nursery 2",
    "NURSERY 3": "Nursery 2",
    "PRIMARY1": "Basic 1",
    "PRIMARY 1": "Basic 1",
    "BASIC1": "Basic 1",
    "PRIMARY2": "Basic 2",
    "PRIMARY 2": "Basic 2",
    "BASIC2": "Basic 2",
    "PRIMARY3": "Basic 3",
    "PRIMARY 3": "Basic 3",
    "BASIC3": "Basic 3",
    "PRIMARY4": "Basic 4",
    "PRIMARY 4": "Basic 4",
    "BASIC4": "Basic 4",
    "PRIMARY5": "Basic 5",
    "PRIMARY 5": "Basic 5",
    "BASIC5": "Basic 5",
    "JSS1": "JSS1",
    "JSS2": "JSS2",
    "JSS3": "JSS3",
    "SS1": "SS1",
    "SS2": "SS2",
    "SS3": "SS3",
    "SSS1": "SS1",
    "SSS2": "SS2",
    "SSS3": "SS3",
    "EXSTUDENT": None,
    "EX-STUDENT": None,
    "EX STUDENT": None,
}

PLACEHOLDER_NAME_RE = re.compile(
    r"\b(test|placeholder|nav test|pwd student|demo student|maxsub)\b",
    re.I,
)


def neat_name(value: str) -> str:
    name = re.sub(r"\s+", " ", (value or "").strip())
    name = name.replace("MAC- PEPPLE", "MAC-PEPPLE")
    if not name:
        return ""
    letters = re.sub(r"[^A-Za-z]", "", name)
    if letters and (letters.isupper() or letters.islower()):
        return name.title()
    return name


def normalize_gender(value: str) -> str:
    raw = (value or "").strip().lower()
    if raw in {"m", "male"}:
        return "Male"
    if raw in {"f", "female"}:
        return "Female"
    return ""


def parse_name_block(raw) -> tuple[str, dict[str, str]]:
    # Spreadsheet sometimes stores dates in the Name column; reject those.
    if raw is None:
        return "", {}
    if hasattr(raw, "isoformat") and not isinstance(raw, str):
        return "", {}
    text = str(raw).replace("\r", "")
    if re.match(r"^\d{4}-\d{2}-\d{2}", text.strip()):
        return "", {}
    lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
    if not lines:
        return "", {}
    name = neat_name(lines[0])
    if not name or re.match(r"^\d{4}-\d{2}-\d{2}", name):
        return "", {}
    meta: dict[str, str] = {}
    for line in lines[1:]:
        if ":" not in line:
            continue
        key, val = line.split(":", 1)
        meta[key.strip().upper()] = val.strip()
    return name, meta


def admission_year_from_regno(regno: str) -> int:
    match = re.match(r"^PCS0?(\d{2})", regno.upper())
    if match:
        return 2000 + int(match.group(1))
    return 2025


def sequence_number_from_regno(regno: str) -> int | None:
    match = re.match(r"^PCS0?\d{2}(\d+)$", regno.upper())
    if not match:
        return None
    try:
        return int(match.group(1))
    except ValueError:
        return None


def clean_phone(value: str) -> str:
    phone = (value or "").strip()
    if phone in {"", "0", "NO PARENT PHONE NUMBER YET"}:
        return ""
    return phone


class Command(BaseCommand):
    help = (
        "Import real students from User List Excel rosters, map classes to "
        "Creche → SS3, assign Arm A, and remove placeholder students."
    )

    def add_arguments(self, parser):
        parser.add_argument(
            "--dir",
            type=str,
            default=str(DEFAULT_ROSTER_DIR),
            help="Directory containing User_List*.xlsx files",
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Parse and report without writing to the database",
        )
        parser.add_argument(
            "--keep-unknown",
            action="store_true",
            help="Do not delete students missing from the roster",
        )

    def handle(self, *args, **options):
        roster_dir = Path(options["dir"])
        if not roster_dir.exists():
            raise CommandError(f"Roster directory not found: {roster_dir}")

        files = sorted(roster_dir.glob("User_List*.xlsx"))
        if not files:
            raise CommandError(f"No User_List*.xlsx files in {roster_dir}")

        rows = self._load_rows(files)
        if not rows:
            raise CommandError("No student rows found in roster files.")

        class_arms = self._class_arm_lookup()
        unknown_classes = sorted(
            {
                row["klass_raw"]
                for row in rows
                if row["klass_raw"] not in CLASS_MAP and row["klass_raw"] != "EXSTUDENT"
            }
        )
        if unknown_classes:
            raise CommandError(f"Unknown class labels in roster: {unknown_classes}")

        created = updated = unchanged = 0
        by_class: dict[str, int] = defaultdict(int)
        ex_count = 0

        with transaction.atomic():
            imported_ids: set[str] = set()
            for row in rows:
                imported_ids.add(row["student_id"])
                level_name = CLASS_MAP[row["klass_raw"]]
                is_active = level_name is not None
                class_arm = class_arms.get((level_name, "A")) if level_name else None
                if level_name and class_arm is None:
                    raise CommandError(
                        f"Missing class arm for {level_name}A — seed class levels first."
                    )

                if options["dry_run"]:
                    if is_active:
                        by_class[level_name] += 1
                    else:
                        ex_count += 1
                    continue

                student, was_created, changed = self._upsert_student(
                    row, class_arm=class_arm, is_active=is_active
                )
                self._link_guardian_parent(student, row)
                if was_created:
                    created += 1
                elif changed:
                    updated += 1
                else:
                    unchanged += 1
                if is_active:
                    by_class[level_name] += 1
                else:
                    ex_count += 1

            removed = 0
            if not options["dry_run"] and not options["keep_unknown"]:
                removed = self._purge_placeholders_and_unknown(imported_ids)

            if not options["dry_run"]:
                self._sync_id_sequences(imported_ids)

        self.stdout.write(self.style.SUCCESS(f"Roster files: {len(files)}"))
        self.stdout.write(f"Rows parsed: {len(rows)}")
        if options["dry_run"]:
            self.stdout.write("Dry run — no database changes.")
        else:
            self.stdout.write(
                self.style.SUCCESS(
                    f"Created {created}, updated {updated}, unchanged {unchanged}, "
                    f"removed placeholders/unknown {removed}"
                )
            )
        self.stdout.write(f"Ex-students: {ex_count}")
        for level in CLASS_LADDER:
            if by_class.get(level):
                self.stdout.write(f"  {level}A: {by_class[level]}")

    def _load_rows(self, files: list[Path]) -> list[dict]:
        by_id: dict[str, dict] = {}
        for path in files:
            wb = openpyxl.load_workbook(path, data_only=True)
            ws = wb.active
            for index, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if index == 1 or not row:
                    continue
                regno = str(row[2] or "").strip().upper()
                if not regno or regno == "REGNO":
                    continue
                name, meta = parse_name_block(row[3])
                if not name:
                    continue
                klass_raw = str(row[4] or "").strip().upper()
                gender = normalize_gender(str(row[5] or ""))
                password = str(row[6] or "school").strip() or "school"
                parent_phone = clean_phone(meta.get("PARENT PHONE", ""))
                parent_name = neat_name(meta.get("PARENT NAME", ""))
                payload = {
                    "student_id": regno,
                    "full_name": name,
                    "klass_raw": klass_raw,
                    "gender": gender,
                    "password": password,
                    "guardian_phone": parent_phone,
                    "guardian_name": parent_name,
                    "admission_year": admission_year_from_regno(regno),
                    "source": path.name,
                }
                # Later files / later rows win on duplicate IDs (none expected).
                by_id[regno] = payload
        # Stable ordering: by class ladder then name.
        level_rank = {label: i for i, label in enumerate(CLASS_LADDER)}

        def sort_key(item: dict):
            level = CLASS_MAP.get(item["klass_raw"])
            rank = level_rank.get(level, 999) if level else 1000
            return (rank, item["full_name"].lower(), item["student_id"])

        return sorted(by_id.values(), key=sort_key)

    def _class_arm_lookup(self) -> dict[tuple[str, str], ClassArm]:
        arms = ClassArm.objects.select_related("class_level")
        return {(arm.class_level.name, arm.name): arm for arm in arms}

    def _upsert_student(
        self, row: dict, *, class_arm: ClassArm | None, is_active: bool
    ) -> tuple[StudentProfile, bool, bool]:
        student_id = row["student_id"]
        student = (
            StudentProfile.objects.select_related("user")
            .filter(student_id__iexact=student_id)
            .first()
        )
        email = f"{student_id.lower()}@students.peaceconceptschool.ng"
        name_parts = row["full_name"].split(" ", 1)
        first_name = name_parts[0]
        last_name = name_parts[1] if len(name_parts) > 1 else ""

        fields = {
            "full_name": row["full_name"],
            "gender": row["gender"],
            "admission_year": row["admission_year"],
            "class_arm": class_arm,
            "is_active": is_active,
            "guardian_name": row["guardian_name"],
            "guardian_phone": row["guardian_phone"],
            "email": email,
        }

        if student is None:
            user = User.objects.create_user(
                email=email,
                password=row["password"],
                username=email,
                account_type=AccountType.STUDENT,
                first_name=first_name,
                last_name=last_name,
                phone=row["guardian_phone"] or "",
                must_change_password=False,
                is_active=is_active,
            )
            student = StudentProfile.objects.create(
                user=user, student_id=student_id, **fields
            )
            return student, True, True

        changed = False
        for key, value in fields.items():
            if getattr(student, key) != value:
                setattr(student, key, value)
                changed = True
        if changed:
            student.save()

        user = student.user
        if user is None:
            user = User.objects.create_user(
                email=email,
                password=row["password"],
                username=email,
                account_type=AccountType.STUDENT,
                first_name=first_name,
                last_name=last_name,
                phone=row["guardian_phone"] or "",
                must_change_password=False,
                is_active=is_active,
            )
            student.user = user
            student.save(update_fields=["user"])
            changed = True
        else:
            user_changed = False
            if user.first_name != first_name:
                user.first_name = first_name
                user_changed = True
            if user.last_name != last_name:
                user.last_name = last_name
                user_changed = True
            if user.email != email:
                user.email = email
                user_changed = True
            if user.username != email:
                # Keep username unique; migrate to email-based username.
                if not User.objects.filter(username=email).exclude(pk=user.pk).exists():
                    user.username = email
                    user_changed = True
            if user.account_type != AccountType.STUDENT:
                user.account_type = AccountType.STUDENT
                user_changed = True
            if user.phone != (row["guardian_phone"] or ""):
                user.phone = row["guardian_phone"] or ""
                user_changed = True
            if user.is_active != is_active:
                user.is_active = is_active
                user_changed = True
            user.set_password(row["password"])
            user.must_change_password = False
            user_changed = True
            if user_changed:
                user.save()
                changed = True

        return student, False, changed

    def _parent_username(self, name: str, phone: str) -> str:
        digits = re.sub(r"\D", "", phone or "")
        if len(digits) >= 7:
            return f"parent.{digits[-10:]}"
        slug = re.sub(r"[^a-z0-9]+", ".", (name or "").lower()).strip(".")
        if slug:
            return f"parent.{slug}"[:40]
        return ""

    def _link_guardian_parent(self, student: StudentProfile, row: dict) -> None:
        name = (row.get("guardian_name") or "").strip()
        phone = (row.get("guardian_phone") or "").strip()
        if not name and not phone:
            return
        username = self._parent_username(name, phone)
        if not username:
            return
        email = f"{username}@parents.peaceconceptschool.ng"
        user = User.objects.filter(username__iexact=username).first()
        if user and user.account_type != AccountType.PARENT:
            username = f"{username}.{student.student_id.lower()}"
            email = f"{username}@parents.peaceconceptschool.ng"
            user = User.objects.filter(username__iexact=username).first()
        if user is None:
            user = User.objects.filter(email__iexact=email).first()
        if user is None:
            display = name or f"Parent of {student.full_name}"
            parts = display.split(" ", 1)
            user = User.objects.create_user(
                email=email,
                password=DEFAULT_PASSWORD,
                username=username,
                account_type=AccountType.PARENT,
                first_name=parts[0],
                last_name=parts[1] if len(parts) > 1 else "",
                phone=phone,
                must_change_password=False,
            )
        parent, _ = ParentProfile.objects.get_or_create(
            user=user,
            defaults={
                "full_name": name or user.get_full_name() or username,
                "phone_number": phone,
            },
        )
        if name and parent.full_name != name:
            parent.full_name = name
            parent.save(update_fields=["full_name"])
        if phone and parent.phone_number != phone:
            parent.phone_number = phone
            parent.save(update_fields=["phone_number"])
        parent.children.add(student)

    def _purge_placeholders_and_unknown(self, imported_ids: set[str]) -> int:
        removed = 0
        qs = StudentProfile.objects.select_related("user").all()
        for student in qs:
            sid = (student.student_id or "").upper()
            name = student.full_name or ""
            unknown = sid not in imported_ids
            placeholder = bool(PLACEHOLDER_NAME_RE.search(name))
            if not (unknown or placeholder):
                continue
            user = student.user
            student.delete()
            if user is not None:
                user.delete()
            removed += 1
        return removed

    def _sync_id_sequences(self, imported_ids: set[str]) -> None:
        by_year: dict[int, int] = defaultdict(int)
        for regno in imported_ids:
            year = admission_year_from_regno(regno)
            number = sequence_number_from_regno(regno)
            if number is not None:
                by_year[year] = max(by_year[year], number)
        for year, last_number in by_year.items():
            seq, _ = StudentIdSequence.objects.get_or_create(
                admission_year=year, defaults={"last_number": last_number}
            )
            if seq.last_number < last_number:
                seq.last_number = last_number
                seq.save(update_fields=["last_number"])
